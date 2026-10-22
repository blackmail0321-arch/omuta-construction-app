import openpyxl,sys,zipfile,datetime
p=sys.argv[1] if len(sys.argv)>1 else '/tmp/omuta-simple.xlsx'
w=openpyxl.load_workbook(p,data_only=False)
assert w.sheetnames==['概要','採用明細','元代価表の構成','採用根拠']
summary={r[0]:r[1] for r in w['概要'].iter_rows(min_row=2,values_only=True)}
assert summary['採用件数']==32;assert summary['採用明細計']==4260400;assert summary['工事費合計']==7209400
rows=list(w['採用明細'].iter_rows(min_row=2,values_only=True));assert len(rows)==32;assert sum(r[6] for r in rows)==4260400
assert all(r[3]*r[5]==r[6] for r in rows)
assert len(w.defined_names)==0
assert len(w._cell_styles)==2
assert w['採用根拠'].max_row==33
assert not any(c.data_type=='f' for s in w for row in s for c in row)
with zipfile.ZipFile(p) as z:
 for info in z.infolist():datetime.datetime(*info.date_time)
 assert len(z.namelist())==9
print('PASS minimal OOXML, 32 line values and sources, 7,209,400 total, valid ZIP timestamps, no legacy names')
